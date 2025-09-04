from io import BytesIO
import mimetypes
import sys
import types
from docx import Document
from fpdf import FPDF
from fastapi.testclient import TestClient


# Provide a stub for the `magic` module used by iocsearcher so tests do not
# depend on the system libmagic library.
magic = types.ModuleType("magic")


def _from_file(path, mime=False):
    if not mime:
        raise NotImplementedError
    mt, _ = mimetypes.guess_type(path)
    return mt or "application/octet-stream"


def _from_buffer(buf, mime=False):
    if not mime:
        raise NotImplementedError
    if buf.startswith(b"%PDF"):
        return "application/pdf"
    if buf.startswith(b"PK"):
        return (
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        )
    if b"<html" in buf.lower():
        return "text/html"
    return "text/plain"


magic.from_file = _from_file
magic.from_buffer = _from_buffer
sys.modules["magic"] = magic

from ioc_checker.main import app

client = TestClient(app)


# Generate sample documents in-memory for upload tests
HTML_BYTES = b"<html>http://example.com</html>"

pdf = FPDF()
pdf.add_page()
pdf.set_font("Arial", size=12)
pdf.cell(200, 10, txt="http://example.com", ln=1)
PDF_BYTES = pdf.output(dest="S").encode("latin1")

doc = Document()
doc.add_paragraph("http://example.com")
buffer = BytesIO()
doc.save(buffer)
DOCX_BYTES = buffer.getvalue()


def test_parse_separates_uri_from_fqdn():
    text = "visit http://example.com and example.org or email test@example.com"
    resp = client.post("/parse", json={"text": text})
    data = resp.json()
    assert "uri" in data and "http://example.com" in data["uri"]
    assert "fqdn" in data and "example.org" in data["fqdn"]
    assert "email" in data and "test@example.com" in data["email"]


def test_parse_file_html():
    resp = client.post(
        "/parse-file",
        files={"file": ("sample.html", BytesIO(HTML_BYTES), "text/html")},
    )
    data = resp.json()
    assert "uri" in data and "http://example.com" in data["uri"]


def test_parse_file_pdf():
    resp = client.post(
        "/parse-file",
        files={"file": ("sample.pdf", BytesIO(PDF_BYTES), "application/pdf")},
    )
    data = resp.json()
    assert "uri" in data and "http://example.com" in data["uri"]


def test_parse_file_docx():
    resp = client.post(
        "/parse-file",
        files={
            "file": (
                "sample.docx",
                BytesIO(DOCX_BYTES),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    data = resp.json()
    assert "uri" in data and "http://example.com" in data["uri"]

